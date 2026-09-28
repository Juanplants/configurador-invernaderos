"""Crea tests/datos/Catalogo_con_errores.xlsx: copia de la plantilla con errores
provocados a propósito, para comprobar que el importador los detecta todos.

Uso: python3 tests/generar_catalogo_con_errores.py
Los errores y avisos esperados están en tests/importacion.js (ESPERADO).
"""
from pathlib import Path
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
ORIGEN = RAIZ / "datos" / "Catalogo_Plantilla_v0.4.xlsx"
DESTINO = RAIZ / "tests" / "datos" / "Catalogo_con_errores.xlsx"

def celda(ws, ident, columna):
    """Celda de la fila cuyo Id es `ident`, en la columna con ese título (fila 4)."""
    col = next(c.column for c in ws[4] if c.value == columna)
    fila = next(c.row for c in ws["A"] if c.value == ident)
    return ws.cell(row=fila, column=col)

wb = load_workbook(ORIGEN)
comp, modelos, perfiles, equipos = wb["Componentes"], wb["Modelos"], wb["Perfiles"], wb["Equipos"]

# --- Errores ---
celda(comp, "C02", "Ref").value = "ARC-42|ARC-99"                  # referencia inexistente
celda(comp, "C05", "Regla").value = "por_metro"                     # regla desconocida
celda(comp, "C09", "Fórmula avanzada").value = "(4+2*ceil(tramos/5))*kg_arriostramento"  # variable inexistente
celda(comp, "C12", "Fórmula avanzada").value = "2*naves*largo+"     # fórmula que no se evalúa
celda(comp, "C04", "Factor").value = "naves*5|naves*6|naves*7"      # 3 variantes para 2 modelos
celda(modelos, "MT-GOT-96", "Alturas a canal admitidas").value = None  # modelo sin alturas
celda(equipos, "PUE-3x3", "Ancho puerta").value = "tres"             # medida de puerta no numérica

# --- Avisos (no impiden cargar) ---
celda(perfiles, "TUB-25", "Precio").value = "consultar"             # precio no numérico
celda(comp, "C07", "Precio unitario").value = None                  # partida sin precio

DESTINO.parent.mkdir(exist_ok=True)
wb.save(DESTINO)
print(DESTINO.relative_to(RAIZ))
