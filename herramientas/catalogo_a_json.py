"""Convierte la plantilla Excel del catálogo en datos/catalogo-ejemplo.json.

Uso: python3 herramientas/catalogo_a_json.py datos/Catalogo_Plantilla_v0.4.xlsx datos/catalogo-ejemplo.json
La app hace la misma conversión en el navegador (js/importador.js, botón
«Cargar catálogo»), y además valida. Si se cambia una, cambiar la otra:
tests/importacion.js comprueba que ambas dan el mismo resultado.
"""
import json, sys, unicodedata, re
from openpyxl import load_workbook

def clave(texto):
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    t = re.sub(r"[^a-zA-Z0-9]+", "_", t).strip("_").lower()
    return t

HOJAS = ["Empresa", "Modelos", "Perfiles", "Componentes", "Cubiertas", "Equipos", "Obra local"]
# Columnas que se guardan como número aunque vengan como texto ("1,4")
NUMERICAS = ("precio", "precio_unitario", "movilizacion", "montaje", "hoyos_y_dados")
NUMERO = re.compile(r"^[-+]?(\d+\.?\d*|\.\d+)(e[-+]?\d+)?$", re.I)

def leer(ws):
    cab = [clave(c.value) if c.value not in (None, "") else None for c in ws[4]]
    filas = []
    for n, fila in enumerate(ws.iter_rows(min_row=6, values_only=True), start=6):
        if not fila or fila[0] in (None, ""):
            continue
        d = {k: v for k, v in zip(cab, fila) if k and v not in (None, "")}
        for k in NUMERICAS:
            v = d.get(k)
            if isinstance(v, str):
                t = v.strip().replace(",", ".", 1)
                if NUMERO.match(t):
                    d[k] = float(t)
                else:
                    print(f"Aviso: {ws.title} fila {n}, {k} = {v!r} no es un número: se trata como sin precio", file=sys.stderr)
                    del d[k]
        filas.append(d)
    return filas

def main(xlsx, salida):
    wb = load_workbook(xlsx, data_only=True)
    cat = {clave(h): leer(wb[h]) for h in HOJAS}
    cat["empresa"] = cat["empresa"][0] if cat["empresa"] else {}
    with open(salida, "w", encoding="utf-8") as f:
        json.dump(cat, f, ensure_ascii=False, indent=2)
    # Copia .js para abrir las páginas con doble clic (file:// no permite leer .json)
    with open(salida.replace(".json", ".js"), "w", encoding="utf-8") as f:
        f.write("// Generado por herramientas/catalogo_a_json.py — no editar a mano\n")
        f.write("window.CATALOGO_EJEMPLO = " + json.dumps(cat, ensure_ascii=False, indent=2) + ";\n")
    print(f"{salida}: " + ", ".join(f"{k} {len(v) if isinstance(v, list) else 1}" for k, v in cat.items()))

if __name__ == "__main__":
    main(*sys.argv[1:3])
